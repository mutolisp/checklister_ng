"""日本名錄匯入：JBIF 維管束植物和名チェックリスト（wamei）+ dao_jp_ylist 合併。

替代 `ylist_import.py`。產出的三個物件結構與舊版完全相同，所以 mobile 端
除了表名之外不需要改動：

  1. jp_names     — 欄位對齊 taicol_names（taicol 專屬欄填 NULL）+ region='JP'。
  2. species_xref — (taxon_id, region, sci_norm, common_name_c)，同 sci_norm
                    跨區 = 共有種。
  3. all_names    — taicol_names ∪ jp_names view。⚠ 只可掃描聚合，絕不可放進
                    JOIN 內側（會物化 ~270k 列、數秒）。

為什麼是「合併」而不是純取代
--------------------------------
JBIF checklist只收維管束植物，且不帶保育／來源屬性。純取代會損失苔蘚類 1,909 筆、
日本特有標記 786、IUCN 1,719、外來/歸化/栽培註記 8,780，以及 915 筆「臺灣產、
JBIF checklist無、YList 有和名」的物種。因此：

  * 和名以 JBIF checklist為準（含同物異名，這是 JBIF checklist最大的價值）
  * 苔蘚類、臺灣種 fallback、保育/來源屬性由 dao_jp_ylist 補

合併鍵 = sci_norm（學名正規化）。**taxon_id 一律沿用該學名在舊 ylist_names 的
id**，只有 JBIF checklist獨有的種才配發新id（y1xxxxxx）——因為使用者的記錄／樣區／
標本都持久化 taxon_id，而解析失敗不會報錯、只會預設顯示空白。

資料來源與授權
--------------
維管束植物和名チェックリスト ver. 1.10（CC BY 4.0）
Yamanouchi, T., Shutoh, K., Osawa, T., Yonekura, K., Kato, S., Shiga, T. 2019.
A checklist of Japanese plant names.
https://gbif.jp/activities/checklist/wamei_checklist_110

CLI（專案根目錄）：
    backend/venv/bin/python -m backend.services.jp_import [db_path]
"""

from __future__ import annotations

import argparse
import sqlite3
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

from backend.db import sqlite_file_path
from backend.utils.backup import backup_db
from backend.services.ylist_import import (
    TAICOL_COLUMNS,
    YLIST_NAME_ID_OFFSET,
    derive_rank,
    load_family_backbone,
    parse_genus,
    sci_norm,
)

WAMEI_XLSX = (
    Path(__file__).resolve().parents[2]
    / "references" / "JP" / "wamei_checklist_ver.1.10.xlsx"
)

#: wamei-only 分類群的 name_id 起點。舊 ylist 佔用 90,000,001..90,020,103，
#: taicol 目前 max 309,598 —— 兩邊都不會撞到這裡。
WAMEI_NAME_ID_OFFSET = 91_000_000

#: wamei-only 分類群的 taxon_id 號段。必須維持 'y' 開頭：mobile 的
#: `isJpTaxonId` 就是 `taxonId.charAt(0) === 'y'`。
WAMEI_TAXON_PREFIX = "y1"

#: 學名來源優先序。GL（Green List）是經同儕審查的正式名錄，故優先；
#: 其餘依覆蓋率 YL > WF > SF。欄位代號見 Hub_data 表頭。
SOURCE_PRIORITY = ("GL", "YL", "WF", "SF")

#: 資料中代表「無此值」的字面值。#N/A 是 Excel 公式殘留，－ 是「無學名」。
_NA = {"#N/A", "", "－", "-", None}

#: wamei 用、但 backbone 以另一個名字收錄的科（同物異名）。
FAMILY_ALIAS = {
    "Asphodelaceae": "Xanthorrhoeaceae",
    "Pentaphylacaceae": "Ternstroemiaceae",
}


def _clean(v) -> str | None:
    """Excel 儲存格 → 乾淨字串，資料中的 #N/A / － 視為 None。"""
    if v is None:
        return None
    s = str(v).strip()
    return None if s in _NA else s


def _deaccent(s: str) -> str:
    """Isoëtaceae → Isoetaceae。backbone 用無變音符號的拼法。"""
    return "".join(
        c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn"
    )


def lookup_family(
    family: str | None, backbone: dict[str, tuple[str, str, str]]
) -> tuple[str, str, str] | None:
    """科名 → (phylum, class, order)，依序嘗試原名、去變音符號、同物異名。"""
    if not family:
        return None
    for key in (family, _deaccent(family), FAMILY_ALIAS.get(family)):
        if key and key in backbone:
            return backbone[key]
    return None


def read_wamei(xlsx_path: Path) -> dict[tuple[str, str], dict]:
    """讀 wamei xlsx，回傳 {(Hub name, lato/stricto): {...}}。

    分類群鍵是 **(Hub name, lato/stricto) 複合鍵**，不是 Hub name —— 同一個
    Hub name 會同時有広義與狹義兩列且學名不同（コスギラン広義 =
    Huperzia selago、狹義 = H. selago var. appressa），只用 Hub name 會把
    1,124 組不同分類群錯誤合併。
    """
    if not xlsx_path.exists():
        raise FileNotFoundError(f"找不到 wamei 檢核表：{xlsx_path}")

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)

    # JN_dataset：種 ID → 學名（不含作者）。同一 ID 有多列（每個和名一列），
    # 取第一個非空的學名即可。
    jn: dict[str, str] = {}
    ws = wb["JN_dataset"]
    for n, row in enumerate(ws.iter_rows(values_only=True)):
        if n == 0:
            continue
        rid, sci = _clean(row[0]), _clean(row[10])
        if rid and sci and rid not in jn:
            jn[rid] = sci

    # Hub_data：和名整合層。
    src_col = {"GL": 6, "SF": 7, "WF": 8, "YL": 9}
    taxa: dict[tuple[str, str], dict] = {}
    ws = wb["Hub_data"]
    for n, row in enumerate(ws.iter_rows(values_only=True)):
        if n == 0:
            continue
        hub = _clean(row[1])
        if not hub:
            continue
        key = (hub, _clean(row[2]) or "")
        t = taxa.setdefault(
            key,
            {
                "hub": hub,
                "lato": _clean(row[2]),
                "names": set(),
                "sci": None,
                "family": None,
                "family_c": None,
                "status": set(),
            },
        )
        all_name = _clean(row[0])
        if all_name:
            t["names"].add(all_name)
        t["family"] = t["family"] or _clean(row[4])
        t["family_c"] = t["family_c"] or _clean(row[5])
        status = _clean(row[10])
        if status:
            t["status"].add(status)
        if t["sci"] is None:
            for code in SOURCE_PRIORITY:
                rid = _clean(row[src_col[code]])
                if rid and jn.get(rid):
                    t["sci"] = jn[rid]
                    break
    wb.close()
    return taxa


def build_rows(
    taxa: dict[tuple[str, str], dict],
    ylist: list[sqlite3.Row],
    backbone: dict[str, tuple[str, str, str]],
) -> tuple[list[dict], dict]:
    """合併 wamei + dao_jp_ylist → jp_names 的列。"""
    audit = {
        "wamei_total": len(taxa),
        "wamei_no_sci": 0,
        "merged": 0,
        "wamei_only": 0,
        "ylist_only": 0,
        "family_no_backbone": set(),
        "hierarchy_null": 0,
        "hybrid": 0,
        "synonyms": 0,
    }

    # 舊 ylist 依 sci_norm 索引 —— taxon_id 要沿用它的，才不會弄壞既有記錄。
    yl_by_sci: dict[str, sqlite3.Row] = {}
    for r in ylist:
        key = sci_norm(r["simple_name"] or "")
        if key:
            yl_by_sci.setdefault(key, r)

    rows: list[dict] = []
    used_sci: set[str] = set()
    #: 已被 wamei 列吸收的舊 taxon_id。剩下的舊列一律原樣保留 —— 舊名錄有
    #: 20,103 列但只有 19,851 個相異學名，若以 sci_norm 去重會丟掉 252 個
    #: 仍被使用者記錄引用的 taxon_id（且連帶失去它們的保育/來源屬性）。
    merged_old_ids: set[str] = set()
    next_wamei_id = 1

    def hierarchy(family: str | None) -> tuple[str | None, str | None, str | None]:
        hit = lookup_family(family, backbone)
        if hit:
            return hit
        if family:
            audit["family_no_backbone"].add(family)
        return (None, None, None)

    # ── 1. wamei 分類群（維管束植物的和名層）──
    for (hub, lato), t in sorted(taxa.items()):
        sci = t["sci"]
        if not sci:
            audit["wamei_no_sci"] += 1
            continue
        key = sci_norm(sci)
        if key in used_sci:
            continue
        used_sci.add(key)

        old = yl_by_sci.get(key)
        if old is not None:
            audit["merged"] += 1
            taxon_id = old["taxon_id"]
            name_id = old["name_id"]
            merged_old_ids.add(taxon_id)
        else:
            audit["wamei_only"] += 1
            taxon_id = f"{WAMEI_TAXON_PREFIX}{next_wamei_id:06d}"
            name_id = WAMEI_NAME_ID_OFFSET + next_wamei_id
            next_wamei_id += 1

        # 同物異名 → alternative_name_c。searchSpeciesJp 已經 LIKE 這一欄，
        # 所以異名搜尋不需要額外的 SQL。
        alts = sorted(n for n in t["names"] if n != hub)
        if alts:
            audit["synonyms"] += len(alts)

        family = t["family"]
        phylum, klass, order = hierarchy(family)
        is_hybrid = "true" if ("×" in sci or " x " in f" {sci} ") else None
        if is_hybrid:
            audit["hybrid"] += 1
        if not all((phylum, klass, order, family, parse_genus(sci))):
            audit["hierarchy_null"] += 1

        rows.append(
            {
                "name_id": name_id,
                "rank": derive_rank(sci),
                "simple_name": sci,
                "name_author": None,
                "formatted_name": sci,
                "usage_status": "accepted",
                "taxon_id": taxon_id,
                "taxon_id_all": taxon_id,
                "is_in_taiwan": None,
                # 和名以 wamei Hub name 為準；広義/狹義 併入顯示名以區分同名
                "common_name_c": f"{hub}{t['lato']}" if t["lato"] else hub,
                "alternative_name_c": ", ".join(alts) if alts else None,
                # 屬性沿用舊 ylist（wamei 不帶保育/來源資訊）
                "is_endemic": (old["is_endemic"] if old is not None else None),
                "alien_type": None,
                "iucn": (old["iucn"] if old is not None else None),
                "redlist": None,
                "kingdom": "Plantae",
                "kingdom_c": "植物界",
                "phylum": phylum,
                "phylum_c": None,
                "class": klass,
                "class_c": None,
                "order": order,
                "order_c": None,
                "family": family,
                "family_c": t["family_c"],
                "genus": parse_genus(sci),
                "genus_c": None,
                "nomenclature_name": None,
                "cites": None,
                "is_fossil": None,
                "is_terrestrial": None,
                "is_freshwater": None,
                "is_brackish": None,
                "is_marine": None,
                "alien_status_note": (
                    old["alien_status_note"] if old is not None else None
                ),
                "protected": None,
                "is_hybrid": is_hybrid,
                "region": "JP",
            }
        )

    # ── 2. 沒被吸收的舊列（苔蘚類、臺灣種 fallback、同學名的重複列）原樣保留 ──
    #    以 taxon_id 而非 sci_norm 判斷，確保每一個舊 id 都還在。
    for r in ylist:
        if r["taxon_id"] in merged_old_ids:
            continue
        audit["ylist_only"] += 1
        rows.append({c: r[c] for c in TAICOL_COLUMNS} | {"region": "JP"})

    return rows, audit


def build(db_path: Path) -> None:
    backbone = load_family_backbone()
    taxa = read_wamei(WAMEI_XLSX)
    backup_path = backup_db(str(db_path))
    print(f"已備份資料庫 → {backup_path}")

    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()

    # 舊表可能還叫 ylist_names（首次遷移）或已經是 jp_names（重跑）。
    src_table = "ylist_names"
    exists = cur.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name='ylist_names';"
    ).fetchone()
    if not exists:
        src_table = "jp_names"
    ylist = cur.execute(
        f'SELECT * FROM "{src_table}";'  # noqa: S608 - 表名來自上面的白名單
    ).fetchall()
    print(f"讀入既有日本名錄 {len(ylist)} 筆（{src_table}）")

    rows, audit = build_rows(taxa, ylist, backbone)

    cols = TAICOL_COLUMNS + ["region"]
    col_defs = ",\n  ".join(
        f'"{c}" INTEGER PRIMARY KEY' if c == "name_id" else f'"{c}" VARCHAR'
        for c in cols
    )
    cur.execute("DROP TABLE IF EXISTS jp_names;")
    cur.execute(f"CREATE TABLE jp_names (\n  {col_defs}\n);")
    placeholders = ", ".join("?" for _ in cols)
    quoted = ", ".join(f'"{c}"' for c in cols)
    cur.executemany(
        f"INSERT INTO jp_names ({quoted}) VALUES ({placeholders});",  # noqa: S608
        [tuple(r[c] for c in cols) for r in rows],
    )
    for col in ("common_name_c", "simple_name", "family", "genus", "taxon_id"):
        cur.execute(f'CREATE INDEX idx_jp_{col} ON jp_names("{col}");')
    cur.execute('CREATE INDEX idx_jp_order ON jp_names("order");')

    # species_xref：TaiCOL accepted + 日本全量，sci_norm 為跨區配對鍵。
    cur.execute("DROP TABLE IF EXISTS species_xref;")
    cur.execute(
        "CREATE TABLE species_xref ("
        "taxon_id TEXT PRIMARY KEY, region TEXT, sci_norm TEXT, common_name_c TEXT);"
    )
    xref = [
        (r["taxon_id"], "TW", sci_norm(r["simple_name"] or ""), r["common_name_c"])
        for r in cur.execute(
            "SELECT taxon_id, simple_name, common_name_c FROM taicol_names "
            "WHERE usage_status='accepted' AND taxon_id IS NOT NULL GROUP BY taxon_id;"
        ).fetchall()
    ]
    xref += [
        (r["taxon_id"], "JP", sci_norm(r["simple_name"] or ""), r["common_name_c"])
        for r in rows
    ]
    cur.executemany(
        "INSERT OR IGNORE INTO species_xref (taxon_id, region, sci_norm, common_name_c) "
        "VALUES (?, ?, ?, ?);",
        xref,
    )
    cur.execute("CREATE INDEX idx_xref_sci_norm ON species_xref(sci_norm);")

    # ⚠ all_names 只給分類樹的「掃描+GROUP BY」用，絕不可放進 JOIN 內側。
    sel = ", ".join(f'"{c}"' for c in TAICOL_COLUMNS)
    cur.execute("DROP VIEW IF EXISTS all_names;")
    cur.execute(
        f"CREATE VIEW all_names AS "
        f"SELECT {sel}, 'TW' AS region FROM taicol_names "
        f"UNION ALL "
        f"SELECT {sel}, region FROM jp_names;"
    )
    conn.commit()

    shared = cur.execute(
        "SELECT count(*) FROM species_xref WHERE region='TW' AND sci_norm IN "
        "(SELECT sci_norm FROM species_xref WHERE region='JP');"
    ).fetchone()[0]
    cur.execute("VACUUM;")
    conn.close()

    print(f"jp_names 匯入 {len(rows)} 筆")
    print(f"  wamei 分類群 {audit['wamei_total']}（無學名而略過 {audit['wamei_no_sci']}）")
    print(f"  沿用舊 taxon_id（兩邊都有）: {audit['merged']}")
    print(f"  wamei 獨有（新配號）        : {audit['wamei_only']}")
    print(f"  舊名錄獨有（苔蘚/臺灣 fallback）: {audit['ylist_only']}")
    print(f"  同物異名                    : {audit['synonyms']}")
    print(f"  雜交種標記                  : {audit['hybrid']}")
    print(f"  階層有缺（分類樹會看不到）    : {audit['hierarchy_null']}")
    if audit["family_no_backbone"]:
        miss = sorted(audit["family_no_backbone"])
        print(f"  ⚠ {len(miss)} 個科無 backbone：{miss[:20]}")
    print(f"  與 TaiCOL 共有種（同 sci_norm）：{shared}")
    print("species_xref + all_names view 已建立")


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(
        description="wamei checklist + dao_jp_ylist → jp_names + species_xref + all_names"
    )
    parser.add_argument(
        "db_path", nargs="?", type=Path, default=Path(sqlite_file_path),
        help=f"twnamelist.db 路徑（預設 {sqlite_file_path}）",
    )
    args = parser.parse_args(argv)
    build(args.db_path)
    return 0


if __name__ == "__main__":
    import sys

    sys.exit(main(sys.argv[1:]))
